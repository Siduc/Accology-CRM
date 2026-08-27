"""Draft Excel working-papers pack from a trial balance already on disk.

Does not pull books, post journals, email, or file. Figures come only from
Current/Source. Credits negative, same convention as the IRIS CSV.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.client import Client
from app.services.client_playbook import (
    client_root_path,
    current_accounts_year,
    ensure_client_pack,
    get_or_create_playbook,
)
from app.services.iris_elements import (
    find_source_tb,
    load_chart_names,
    map_name,
    parse_tb_file,
)

NAVY = "052891"
HEADER = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
THIN = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)


@dataclass
class PackResult:
    ok: bool
    error: str = ""
    source_file: str = ""
    pack_file: str = ""
    queries_file: str = ""
    rows: int = 0
    mapped: int = 0
    net: float = 0.0
    queries: List[str] = field(default_factory=list)


def _money_cell(ws, r: int, c: int, value: float) -> None:
    cell = ws.cell(r, c, round(float(value), 2))
    cell.number_format = '#,##0.00;(#,##0.00);"-"'
    cell.border = THIN


def write_draft_pack(
    dest: Path,
    *,
    company: str,
    as_at: date,
    source_name: str,
    lines: List[Dict[str, Any]],
    queries: List[str],
) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    ye = as_at.strftime("%d %B %Y")
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = company
    cover["A1"].font = Font(name="Calibri", bold=True, size=18, color=NAVY)
    cover["A2"] = f"Draft working papers — year ended {ye}"
    cover["A3"] = f"Source file: {source_name}"
    cover["A4"] = (
        "DRAFT only. Not an IRIS statutory set. Do not file. "
        "Credits are negative. Nothing invented — remap of the source trial balance."
    )
    cover["A4"].font = Font(name="Calibri", italic=True, color="833C0C")
    cover["A6"] = "Open queries"
    cover["A6"].font = Font(bold=True, color=NAVY)
    if queries:
        for i, q in enumerate(queries, 1):
            cover.cell(6 + i, 1, f"{i}. {q}")
    else:
        cover["A7"] = "None raised from this TB remap."
    cover.column_dimensions["A"].width = 110

    qws = wb.create_sheet("Queries")
    qws["A1"] = "No"
    qws["B1"] = "Query"
    qws["C1"] = "Before IRIS final?"
    for col in range(1, 4):
        qws.cell(1, col).fill = HEADER
        qws.cell(1, col).font = HEADER_FONT
    if queries:
        for i, q in enumerate(queries, 1):
            qws.cell(i + 1, 1, i)
            qws.cell(i + 1, 2, q)
            qws.cell(i + 1, 2).alignment = Alignment(wrap_text=True)
            qws.cell(i + 1, 3, "Yes — confirm before filing")
    else:
        qws["A2"] = "None"
    qws.column_dimensions["A"].width = 6
    qws.column_dimensions["B"].width = 100
    qws.column_dimensions["C"].width = 28

    tb = wb.create_sheet("Trial balance")
    headers = ["Source account", "Accology / IRIS name", "Signed amount", "Map"]
    for i, h in enumerate(headers, 1):
        cell = tb.cell(1, i, h)
        cell.fill = HEADER
        cell.font = HEADER_FONT
    net = 0.0
    for r, line in enumerate(lines, 2):
        tb.cell(r, 1, line["source"]).border = THIN
        tb.cell(r, 2, line["iris"]).border = THIN
        _money_cell(tb, r, 3, line["amount"])
        tb.cell(r, 4, line["map"]).border = THIN
        net += line["amount"]
    last = len(lines) + 2
    tb.cell(last, 1, "Net (must be 0.00)").font = Font(bold=True)
    _money_cell(tb, last, 3, net)
    for col, w in enumerate([48, 56, 18, 14], 1):
        tb.column_dimensions[get_column_letter(col)].width = w

    wb.save(dest)
    wb.close()


def build_client_pack(db: Session, client: Client) -> PackResult:
    pb = get_or_create_playbook(db, client.id)
    root = client_root_path(client)
    if not root:
        return PackResult(ok=False, error="No client folder on this laptop’s OneDrive.")
    ensure_client_pack(db, client)
    source_dir = root / "Current" / "Source"
    wp = root / "Current" / "Working Papers"
    wp.mkdir(parents=True, exist_ok=True)

    names = load_chart_names()
    src = find_source_tb(source_dir)
    if not src or not src.is_file():
        return PackResult(
            ok=False,
            error="No trial balance in Current/Source. Pull books or drop a TB CSV first.",
        )
    raw = parse_tb_file(src)
    if not raw:
        return PackResult(
            ok=False,
            error=f"Could not read account/amount columns from {src.name}.",
            source_file=str(src),
        )

    from calendar import monthrange

    year = current_accounts_year(db, client, pb)
    month = min(12, max(1, int(pb.year_end_month or 12)))
    day = int(pb.year_end_day or 31)
    last = monthrange(year, month)[1]
    as_at = date(year, month, min(max(1, day), last))

    lines: List[Dict[str, Any]] = []
    queries: List[str] = []
    mapped = 0
    net = 0.0
    for row in raw:
        dest, score = map_name(row["source"], names) if names else (row["source"], 0.0)
        how = "mapped" if names and dest in names else "unmapped"
        if how == "mapped":
            mapped += 1
        else:
            queries.append(f"Unmapped source account {row['source']!r} — confirm Accology Chart name.")
        if names and dest in names and score < 0.86:
            queries.append(
                f"Soft map {row['source']!r} → {dest!r} (score {score:.2f}) — confirm wording."
            )
        amt = round(float(row["amount"]), 2)
        net = round(net + amt, 2)
        lines.append({"source": row["source"], "iris": dest, "amount": amt, "map": how})
    if abs(net) >= 0.005:
        queries.append(f"Trial balance net is {net:.2f}, not 0.00 — do not import to IRIS until this is explained.")

    ye_words = as_at.strftime("%d %B %Y")
    safe = (client.company_name or "Client").strip()
    pack_path = wp / f"{safe} - {ye_words} - Draft working papers.xlsx"
    q_path = wp / f"{safe} - {ye_words} - Queries.md"
    write_draft_pack(
        pack_path,
        company=client.company_name or "",
        as_at=as_at,
        source_name=src.name,
        lines=lines,
        queries=queries,
    )
    q_path.write_text(
        f"# Queries — {safe} year ended {ye_words}\n\n"
        f"Source: `{src.name}`\n\n"
        + ("\n".join(f"- {q}" for q in queries) if queries else "- None raised.\n"),
        encoding="utf-8",
    )
    return PackResult(
        ok=True,
        source_file=str(src),
        pack_file=str(pack_path),
        queries_file=str(q_path),
        rows=len(lines),
        mapped=mapped,
        net=net,
        queries=queries,
    )
